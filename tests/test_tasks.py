import asyncio
import json

import pytest

from tengyu.files import scan_directory
from tengyu.models import GenerationResult, TitleOptions
from tengyu.providers import ProviderError
from tengyu.storage import Store
from tengyu.tasks import TaskManager


class FakeProvider:
    def __init__(self, error=None, delay=0):
        self.calls = 0
        self.error = error
        self.delay = delay

    async def generate_title(self, *args):
        self.calls += 1
        await asyncio.sleep(self.delay)
        if self.error:
            raise self.error
        return GenerationResult(raw=json.dumps({"title": "Blue product", "char_count": 12}),
                                finish="stop", elapsed=0.1)


def setup(tmp_path, make_image, credentials, provider, count=3, **options):
    images = tmp_path / "商品目录"
    for i in range(count):
        make_image(images / f"{i:03d}.jpg")
    store = Store(tmp_path / "state")
    scan = scan_directory(str(images), "images")
    store.put("scan", scan["id"], scan)
    manager = TaskManager(store, credentials, {"aliyun": provider})
    batch = manager.create(scan["id"], TitleOptions(**options))
    return manager, batch["id"]


async def test_single_generation_review_and_selected_retry(tmp_path, make_image, credentials):
    provider = FakeProvider()
    manager, batch = setup(tmp_path, make_image, credentials, provider, min_chars=30)
    await manager.start(batch)
    await manager.tasks[batch]
    items = manager.store.items(batch)
    assert provider.calls == 3 and all(i["status"] == "review" for i in items)
    assert manager.summary(batch)["exported_at"]
    assert all(i["title"] == "Blue product" and i["reported_char_count"] == 12 for i in items)
    attempt = manager.store.attempts(items[0]["id"])[0]
    assert attempt["char_count"] == attempt["reported_char_count"] == 12
    from openpyxl import load_workbook
    book = load_workbook(manager.summary(batch)["excel_path"])
    assert list(book.active.values)[1] == ("000", "Blue product")
    assert book.active.max_column == 2
    book.close()
    await manager.retry(batch, [items[1]["id"]])
    await manager.tasks[batch]
    assert provider.calls == 4
    assert len(manager.store.attempts(items[1]["id"])) == 2
    assert len(manager.store.attempts(items[0]["id"])) == 1


async def test_pause_stops_dispatch_and_retry_does_not_run_unselected_pending(tmp_path, make_image, credentials):
    provider = FakeProvider(delay=0.05)
    manager, batch = setup(tmp_path, make_image, credentials, provider, concurrency=1)
    await manager.start(batch)
    while provider.calls == 0:
        await asyncio.sleep(0.005)
    manager.pause(batch)
    await manager.tasks[batch]
    assert provider.calls == 1
    item = manager.store.items(batch)[0]
    await manager.retry(batch, [item["id"]])
    await manager.tasks[batch]
    assert provider.calls == 2 and manager.summary(batch)["counts"]["pending"] == 2
    await manager.start(batch)
    await manager.tasks[batch]
    assert provider.calls == 4


async def test_global_error_stops_dispatch(tmp_path, make_image, credentials):
    provider = FakeProvider(ProviderError("余额不足", global_error=True))
    manager, batch = setup(tmp_path, make_image, credentials, provider, concurrency=1)
    await manager.start(batch)
    await manager.tasks[batch]
    assert provider.calls == 1
    assert manager.summary(batch)["status"] == "paused"
    assert manager.summary(batch)["counts"]["pending"] == 2


async def test_export_failure_never_generates_again(tmp_path, make_image, credentials, monkeypatch):
    provider = FakeProvider()
    manager, batch = setup(tmp_path, make_image, credentials, provider)
    import tengyu.tasks as tasks
    original = tasks.export_excel
    def fail(*args):
        raise ValueError("文件被占用")
    monkeypatch.setattr(tasks, "export_excel", fail)
    await manager.start(batch)
    await manager.tasks[batch]
    assert manager.summary(batch)["export_error"] == "文件被占用"
    monkeypatch.setattr(tasks, "export_excel", original)
    await manager.export(batch)
    assert provider.calls == 3 and manager.summary(batch)["exported_at"]


def test_recovery_retains_completed_and_marks_running_unknown(tmp_path, make_image, credentials):
    manager, batch = setup(tmp_path, make_image, credentials, FakeProvider())
    items = manager.store.items(batch)
    items[0].update(status="completed", raw="kept", title="kept")
    items[1].update(status="running")
    manager.store.save_item(items[0])
    manager.store.save_item(items[1])
    attempt = manager.store.begin_attempt(items[1]["id"])
    record = manager.store.get("batch", batch)
    record["status"] = "running"
    manager.store.put("batch", batch, record)
    manager.store.recover()
    assert manager.store.items(batch)[0]["title"] == "kept"
    assert manager.store.items(batch)[1]["status"] == "unknown"
    assert manager.store.attempts(items[1]["id"])[0]["id"] == attempt
    assert manager.store.attempts(items[1]["id"])[0]["status"] == "unknown"


async def test_changed_image_not_submitted(tmp_path, make_image, credentials):
    provider = FakeProvider()
    manager, batch = setup(tmp_path, make_image, credentials, provider, count=1)
    item = manager.store.items(batch)[0]
    from pathlib import Path
    make_image(Path(item["image_path"]), "green")
    await manager.start(batch)
    await manager.tasks[batch]
    assert provider.calls == 0 and manager.summary(batch)["counts"]["failed"] == 1


async def test_duplicate_start_is_rejected(tmp_path, make_image, credentials):
    provider = FakeProvider(delay=0.03)
    manager, batch = setup(tmp_path, make_image, credentials, provider)
    await manager.start(batch)
    with pytest.raises(ValueError):
        await manager.start(batch)
    await manager.tasks[batch]
    assert provider.calls == 3


async def test_legacy_batch_keeps_plain_text_protocol(tmp_path, make_image, credentials):
    class LegacyProvider:
        async def generate_title(self, settings, key, options, skill, image):
            assert skill == "最终只输出标题本身。"
            return GenerationResult(raw="Blue product", finish="stop")

    manager, batch = setup(tmp_path, make_image, credentials, LegacyProvider(), count=1)
    record = manager.store.get("batch", batch)
    record.pop("output_format")
    record["skill"] = "最终只输出标题本身。"
    manager.store.put("batch", batch, record)
    await manager.start(batch)
    await manager.tasks[batch]
    item = manager.store.items(batch)[0]
    assert item["status"] == "completed" and item["title"] == "Blue product"
    assert item["reported_char_count"] is None


async def test_retry_clears_previous_reported_count(tmp_path, make_image, credentials):
    provider = FakeProvider()
    manager, batch = setup(tmp_path, make_image, credentials, provider, count=1)
    await manager.start(batch)
    await manager.tasks[batch]
    item = manager.store.items(batch)[0]
    provider.error = ProviderError("请求失败")
    await manager.retry(batch, [item["id"]])
    await manager.tasks[batch]
    item = manager.store.item(item["id"])
    assert item["reported_char_count"] is None and item["char_count"] is None
    assert item["title"] is None and item["status"] == "failed"

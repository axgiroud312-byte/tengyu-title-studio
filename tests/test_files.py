from pathlib import Path

import pytest
from openpyxl import load_workbook

from tengyu.files import encoded_image, evaluate_title, export_excel, scan_directory, verify_source


def test_sku_natural_order_and_first_corrupt(tmp_path, make_image):
    for sku in ["10", "2", "001"]:
        make_image(tmp_path / sku / "10.jpg")
        make_image(tmp_path / sku / "2.jpg")
    (tmp_path / "4-empty").mkdir()
    (tmp_path / "5-broken").mkdir()
    (tmp_path / "5-broken" / "1.jpg").write_bytes(b"broken")
    make_image(tmp_path / "5-broken" / "2.jpg")
    make_image(tmp_path / "6-nested" / "more" / "1.jpg")
    make_image(tmp_path / ".hidden" / "1.jpg")
    result = scan_directory(str(tmp_path), "sku")
    assert [i["sku"] for i in result["items"]] == ["001", "2", "4-empty", "5-broken", "6-nested", "10"]
    assert Path(result["items"][0]["image_path"]).name == "2.jpg"
    assert [i["status"] for i in result["items"]] == ["pending", "pending", "failed", "failed", "failed", "pending"]
    assert result["items"][3]["image_path"].endswith("1.jpg")


def test_images_suffix_duplicates_and_nonrecursive(tmp_path, make_image):
    make_image(tmp_path / "001.front.jpg")
    make_image(tmp_path / "001.front.png")
    make_image(tmp_path / "more" / "002.jpg")
    (tmp_path / "标题.xlsx").write_bytes(b"ignored")
    result = scan_directory(str(tmp_path), "images")
    assert [i["sku"] for i in result["items"]] == ["001.front", "001.front"]
    assert all(i["warnings"] for i in result["items"])


def test_image_changes_are_not_silently_submitted(tmp_path, make_image):
    path = make_image(tmp_path / "1.jpg")
    item = scan_directory(str(tmp_path), "images")["items"][0]
    original = path.read_bytes()
    assert encoded_image(item).startswith("data:image/jpeg;base64,")
    assert path.read_bytes() == original
    make_image(path, "red")
    with pytest.raises(ValueError, match="变化"):
        verify_source(item)


@pytest.mark.parametrize("raw,count", [(" 商品  title ", 8), ("👩‍👩‍👧‍👦", 1), ("e\u0301", 1), ("ก้", 1), ("ABC", 3)])
def test_visible_characters(raw, count):
    result = evaluate_title(raw, "stop", count, count)
    assert result["char_count"] == count
    assert result["status"] == "completed"


@pytest.mark.parametrize("raw,finish", [("", "stop"), ("hello", "length"), ('{"title":"ABC"}', "stop"), ("one\ntwo", "stop"), ("标题：abc", "stop")])
def test_incomplete_or_malformed_is_not_a_title(raw, finish):
    result = evaluate_title(raw, finish, None, None)
    assert result["status"] == "failed"
    assert result["title"] is None


def test_bad_length_preserved():
    result = evaluate_title("ABC", "stop", 5, 10)
    assert result["status"] == "review"
    assert result["title"] == "ABC"


def test_excel_exact_two_columns_literal_and_warning(tmp_path):
    path = tmp_path / "标题.xlsx"
    rows = [
        {"sku": "00123", "title": "=SUM(1,2)", "status": "completed"},
        {"sku": "+002", "title": "short", "status": "review", "error": "长度不足"},
        {"sku": "003", "title": "stale title", "status": "failed"},
    ]
    export_excel({"excel_path": str(path)}, rows)
    book = load_workbook(path)
    sheet = book.active
    assert list(sheet.values) == [("货号", "标题"), ("00123", "=SUM(1,2)"), ("+002", "short"), ("003", None)]
    assert sheet.max_column == 2 and len(book.worksheets) == 1
    assert sheet["A2"].data_type == "s" and sheet["B2"].data_type == "s"
    assert sheet["B3"].comment.text == "长度不足"
    book.close()


def test_export_preserves_existing_file_on_invalid_data(tmp_path):
    path = tmp_path / "标题.xlsx"
    path.write_bytes(b"previous")
    with pytest.raises(ValueError):
        export_excel({"excel_path": str(path)}, [{"sku": "1", "title": "bad\x00text", "status": "completed"}])
    assert path.read_bytes() == b"previous"

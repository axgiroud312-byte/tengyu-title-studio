from __future__ import annotations

import base64
import hashlib
import io
import os
import re
import tempfile
import uuid
from collections import Counter
from pathlib import Path

import regex
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps, UnidentifiedImageError

EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}


def natural_key(path: Path):
    return ([(1, int(p)) if p.isdigit() else (0, p.casefold())
             for p in re.split(r"(\d+)", path.name)], path.name)


def visible(path: Path) -> bool:
    if path.name.startswith(".") or path.is_symlink():
        return False
    return not bool(getattr(path.stat(), "st_file_attributes", 0) & 0x6)


def fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_image(path: Path) -> Image.Image:
    if not path.is_file() or path.is_symlink():
        raise ValueError("图片不存在或是快捷链接，请重新扫描商品文件夹。")
    if path.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("图片超过 100 MB，请先缩小图片后重新扫描。")
    try:
        with Image.open(path) as image:
            if image.width * image.height > 40_000_000:
                raise ValueError("图片超过 4000 万像素，请先缩小图片。")
            image.seek(0)
            image.load()
            image = ImageOps.exif_transpose(image)
            if "A" in image.getbands() or "transparency" in image.info:
                rgba = image.convert("RGBA")
                background = Image.new("RGB", rgba.size, "white")
                background.paste(rgba, mask=rgba.getchannel("A"))
                return background
            return image.convert("RGB")
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise ValueError("图片损坏或无法读取，请更换图片后重新扫描。") from exc


def check_output(directory: Path):
    if not directory.is_dir():
        raise ValueError("商品文件夹不存在，请重新选择。")
    try:
        with tempfile.TemporaryFile(dir=directory):
            pass
    except OSError as exc:
        raise ValueError("无法在商品文件夹保存 Excel，请检查文件夹写入权限和磁盘空间。") from exc


def scan_directory(directory: str, mode: str) -> dict:
    root = Path(directory).expanduser().resolve(strict=True)
    check_output(root)
    children = sorted((p for p in root.iterdir() if visible(p)), key=natural_key)
    sources = [p for p in children if p.is_dir()] if mode == "sku" else [
        p for p in children if p.is_file() and p.suffix.lower() in EXTENSIONS]
    items = []
    for position, source in enumerate(sources):
        sku = source.name if mode == "sku" else source.stem
        path, error, checksum = None, None, None
        try:
            if mode == "sku":
                images = sorted((p for p in source.iterdir() if visible(p) and p.is_file()
                                 and p.suffix.lower() in EXTENSIONS), key=natural_key)
                path = images[0] if images else None
            else:
                path = source
            if not path:
                raise ValueError("此货号文件夹内没有支持的图片。")
            read_image(path).close()
            checksum = fingerprint(path)
        except (ValueError, OSError) as exc:
            error = str(exc) if isinstance(exc, ValueError) else "无法读取文件，请检查文件夹访问权限。"
        items.append({"id": uuid.uuid4().hex, "position": position, "sku": sku,
                      "image_path": str(path) if path else None, "fingerprint": checksum,
                      "status": "failed" if error else "pending", "input_error": error,
                      "error": error, "warnings": [], "title": None, "raw": None,
                      "char_count": None, "elapsed": None})
    counts = Counter(i["sku"].casefold() for i in items)
    for item in items:
        if counts[item["sku"].casefold()] > 1:
            item["warnings"].append("货号重复：各图片仍分别生成一行。")
    return {"id": uuid.uuid4().hex, "mode": mode, "directory": str(root), "items": items}


def verify_source(item: dict):
    path = Path(item["image_path"] or "")
    if not path.is_file() or path.is_symlink() or fingerprint(path) != item["fingerprint"]:
        raise ValueError("源图片已变化或被移走，请重新扫描；此次没有调用模型。")


def encoded_image(item: dict) -> str:
    verify_source(item)
    image = read_image(Path(item["image_path"]))
    try:
        image.thumbnail((2048, 2048))
        out = io.BytesIO()
        image.save(out, "JPEG", quality=88, optimize=True)
        encoded = base64.b64encode(out.getvalue()).decode("ascii")
        if len(encoded) > 9_000_000:
            raise ValueError("压缩后的图片仍超出接口限制，请缩小原图片后重新扫描。")
        return "data:image/jpeg;base64," + encoded
    finally:
        image.close()


def thumbnail(item: dict) -> bytes:
    verify_source(item)
    image = read_image(Path(item["image_path"]))
    try:
        image.thumbnail((640, 640))
        out = io.BytesIO()
        image.save(out, "JPEG", quality=82)
        return out.getvalue()
    finally:
        image.close()


def evaluate_title(raw: str, finish: str | None, minimum: int | None, maximum: int | None) -> dict:
    # Original body is stored separately. No language/semantic claim is inferred here.
    if finish != "stop":
        return {"status": "failed", "title": None, "char_count": None,
                "error": "模型响应未完整结束，原始内容已保存。"}
    text = raw.strip()
    if not text:
        return {"status": "failed", "title": None, "char_count": None, "error": "模型未返回标题正文。"}
    if ("\n" in text or "\r" in text or text.startswith(("{", "[", "```", "标题：", "标题:", '"', "“"))
            or re.match(r"^(?:\d+[.)、]|[-*#])\s", text) or "<think>" in text.lower()):
        return {"status": "failed", "title": None, "char_count": None,
                "error": "返回内容不是单条标题，请查看原始结果后决定是否重新生成。"}
    title = regex.sub(r"\s+", " ", text)
    count = len(regex.findall(r"\X", title))
    issues = []
    if minimum and count < minimum:
        issues.append(f"字符数为 {count}，少于要求的 {minimum}")
    if maximum and count > maximum:
        issues.append(f"字符数为 {count}，超过要求的 {maximum}")
    if "|" in title:
        issues.append("含标题规则禁用的竖线符号")
    return {"title": title, "char_count": count, "status": "review" if issues else "completed",
            "error": "；".join(issues) or None}


def export_excel(batch: dict, items: list[dict]) -> str:
    path = Path(batch["excel_path"])
    check_output(path.parent)
    book = Workbook()
    sheet = book.active
    sheet.title = "商品标题"
    sheet.append(["货号", "标题"])
    for item in items:
        # Failed/unknown retries do not export stale earlier titles.
        title = item["title"] if item["status"] in {"completed", "review"} else ""
        values = [item["sku"], title or ""]
        if any(len(v) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", v) for v in values):
            book.close()
            raise ValueError("结果含 Excel 不支持的内容，原始结果已保存，未截断或重新生成。")
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if item["status"] == "review":
            cell = sheet.cell(sheet.max_row, 2)
            cell.fill = PatternFill("solid", fgColor="FFF1C2")
            cell.comment = Comment(item["error"] or "请检查此标题", "腾域 ai 写标题")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 95
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    temp = path.with_name(f".{path.stem}-{uuid.uuid4().hex}.tmp.xlsx")
    try:
        book.save(temp)
        os.replace(temp, path)
    except PermissionError as exc:
        raise ValueError("Excel 保存失败，请关闭已打开的表格并检查文件夹写入权限，再点击‘重新导出 Excel’。") from exc
    finally:
        book.close()
        temp.unlink(missing_ok=True)
    return str(path)
